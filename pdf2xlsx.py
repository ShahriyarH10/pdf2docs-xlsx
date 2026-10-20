#!/usr/bin/env python3
"""
pdf2xlsx.py - scanned Bangla/Arabic book (PDF)  ->  one Excel workbook (chapters / sections / hadiths) + Word file.

What it does
  1. Extracts the original page JPEGs from the PDF (no re-rendering), named page_001.jpg ...
  2. OCRs every page with Tesseract (ara+ben), saving one JSON per page (resumable).
  3. Classifies every OCR line: header / noise / arabic / heading / footnote / body.
  4. Joins the lines of the WHOLE book and splits it into hadiths at "[number]" markers
     (misread numbers are repaired from the running sequence, suspicious rows are reported).
  5. Reads the table of contents (leader dots removed before OCR) for the chapter ("অধ্যায়") and section lists.
  6. Writes book.xlsx - ONE workbook with the sheets chapters (id, name), sections (section_id, name) and
     hadiths (id, hadith), Kalpurush, justified - plus book.docx (chapter / section headings + numbered hadiths, Kalpurush, justified),
     review.csv (flagged hadiths) and optional debug images.

Setup
  pip install pymupdf pytesseract pillow openpyxl python-docx numpy opencv-python-headless
  Tesseract with Bangla + Arabic models:
    Ubuntu/Debian: sudo apt install tesseract-ocr tesseract-ocr-ben tesseract-ocr-ara
    macOS:         brew install tesseract tesseract-lang
    Windows:       winget install UB-Mannheim.TesseractOCR   (found automatically in C:\\Program Files\\Tesseract-OCR;
                   with the tessdata_best folder next to this script no extra language data is needed)
  Better accuracy: put the "tessdata_best" ben.traineddata + ara.traineddata (github.com/tesseract-ocr/tessdata_best)
  in a folder called tessdata_best next to this script (used automatically), or pass --tessdata <folder>.
  chapters.txt next to this script (one name per line) overrides the OCR'd chapter names: the bold chapter lines
  of the table of contents are the one thing Tesseract reads badly.

Usage
  python pdf2xlsx.py                                    # a dialog asks for the PDF; results go to the folder
                                                         # "<PDF name> - output" next to it
  python pdf2xlsx.py book.pdf                           # same, PDF given directly (resumes if interrupted)
  python pdf2xlsx.py book.pdf --pages 100-109 --debug   # small test first
  python pdf2xlsx.py book.pdf --rebuild-only            # redo everything except the OCR (seconds), after tuning
  Every run writes book.xlsx (3 sheets), book.docx and review.csv (--no-docx skips the Word file).
  The table of contents is found automatically; a PDF without one (an excerpt) gets its chapters and sections
  from the headings in the text. --toc-pages 17-25 / --toc-pages none override that; --title sets the docx title.

Tuning knobs (if the debug images show misclassified lines): see the CONFIG block below.
The Kalpurush font must be installed on the computer that OPENS the xlsx files (Excel falls back to another font otherwise).
"""
import argparse
import csv
import json
import os
import re
import statistics
import sys
import time
from collections import defaultdict

# Tesseract is already multi-process here; stop it from also multi-threading per page.
os.environ.setdefault("OMP_THREAD_LIMIT", "1")

# ----------------------------------------------------------------------------- CONFIG
CONFIG = dict(
    header_zone=0.075,        # lines whose centre is above this fraction of page height = running header
    ornament_zone=0.95,       # lines below this fraction, if tiny/low-confidence = bottom ornament
    footnote_zone=0.72,       # footnotes can only start below this fraction of page height
    footnote_height_ratio=0.82,   # footnote line is shorter than this x median body line height
    arabic_ratio=0.50,        # line is Arabic if Arabic chars / (Arabic+Bangla chars) >= this
    arabic_ink_ratio=0.75,    # ... or if it is lighter than this x body ink density ...
    arabic_ratio_if_light=0.10,   # ... and at least this Arabic share (rescues badly-OCRed Arabic lines)
    heading_ink_ratio=1.70,   # heading: bold, i.e. ink density above this x median body ink density
    heading_huge_ratio=1.60,  # ... or line taller than this x median body line height
    heading_max_chars=110,
    heading_center_tol=0.15,  # heading must be full-width, or centred within this fraction of page width
    low_conf=60,              # entry flagged if mean OCR confidence is below this
    min_text_chars=15,        # entry flagged if Bangla text shorter than this
    struct_head_ink=1.2,      # layout-detected heading (centred, indented both sides) must be this x bolder
    toc_head_ink=1.25,        # a body line must be at least this x bolder (ink) ...
    toc_head_h=1.15,          # ... or this x taller than the text to be tested against the TOC titles
    font="Kalpurush",         # font of the Bangla cells in the output workbooks
)
# ------------------------------------------------------------------------------------

AR_RE = re.compile(r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]")
BN_RE = re.compile(r"[\u0980-\u09FF]")
BN_DIGITS = "০১২৩৪৫৬৭৮৯"
TO_ASCII = str.maketrans(BN_DIGITS, "0123456789")
TO_BN = str.maketrans("0123456789", BN_DIGITS)
# hadith marker at the start of a line: "[১৪৪]"  "[১৭২।"  "২৯৬]" (opener lost)  "و [২০৯]" (Arabic speck in front).
# NOT "(২)" - that is a numbered list inside a hadith.  group(1) = number with opener, group(2) = without.
MARKER_RE = re.compile(r"^[\s\u0600-\u06FF\u200e\u200f]*(?:[\[\{【\|]\s*([০-৯0-9]{1,4})\s*[\]\)\}】\|।]|([০-৯0-9]{1,4})\s*[\]\}】])\s*")
EMBED_RE = re.compile(r"[\[\{【]\s*([০-৯0-9]{1,4})\s*[\]\)\}】\|।]\s*")      # a marker in the MIDDLE of a line
FOOT_START_RE = re.compile(r"^\s*\W{0,3}([০-৯0-9]{1,4})\s*[\.।:]\s*")


# OCR misreadings that are not Bangla words at all, so replacing them is safe (bold type confuses ভ/ত, ু/ূ, ন্ন/র)
COMMON_FIXES = [
    (r"সম্[তভ]ান", "সন্তান"), (r"তালো", "ভালো"), (r"অতিশপ্ত", "অভিশপ্ত"), (r"জারাত", "জান্নাত"),
    (r"অমূসলিম", "অমুসলিম"), (r"দূআ", "দুআ"), (r"ঘে ব্যক্তি", "যে ব্যক্তি"), (r"সূরে", "সুরে"),
    (r"খিষ্টান", "খ্রিষ্টান"), (r"স্বর[ীা]প", "স্বরূপ"), (r"মসলমান", "মুসলমান"),
]


def common_fixes(text):
    for pat, rep_ in COMMON_FIXES:
        text = re.sub(pat, rep_, text)
    return text


def bn_int(s):
    try:
        return int(s.translate(TO_ASCII))
    except ValueError:
        return None


# =============================================================================== 1. images
def parse_pages(spec, n):
    if not spec:
        return list(range(1, n + 1))
    out = []
    for part in spec.split(","):
        if "-" in part:
            a, b = part.split("-")
            out += range(int(a), min(int(b), n) + 1)
        else:
            out.append(int(part))
    return sorted({p for p in out if 1 <= p <= n})


def extract_images(pdf, img_dir, pages):
    import pymupdf
    os.makedirs(img_dir, exist_ok=True)
    doc = pymupdf.open(pdf)
    paths = {}
    for p in parse_pages(pages, len(doc)) if isinstance(pages, str) or pages is None else pages:
        existing = [f for f in os.listdir(img_dir) if f.startswith(f"page_{p:03d}.")]
        if existing:
            paths[p] = os.path.join(img_dir, existing[0])
            continue
        page = doc[p - 1]
        imgs = page.get_images(full=True)
        if imgs:  # scanned page: take the largest embedded image untouched
            best = max(imgs, key=lambda i: i[2] * i[3])
            info = doc.extract_image(best[0])
            ext = "jpg" if info["ext"] in ("jpeg", "jpg") else info["ext"]
            path = os.path.join(img_dir, f"page_{p:03d}.{ext}")
            with open(path, "wb") as f:
                f.write(info["image"])
        else:     # digital page: render it
            path = os.path.join(img_dir, f"page_{p:03d}.png")
            page.get_pixmap(dpi=300).save(path)
        paths[p] = path
    return paths


# =============================================================================== 2. OCR
def _fix_tail(img, d, idx, tess_cmd, tessdata):
    """A footnote superscript glued to the last word often wrecks it: 'আরো দান করো।⁶³' comes out as 'আরো দান 37'.
    The last token of a Bangla line is re-read from a crop without its right-hand 20% (where the superscript sits)."""
    import pytesseract
    from PIL import ImageOps
    if len(idx) < 3 or any(AR_RE.search(d["text"][i]) for i in idx[:-1]):
        return
    last = idx[-1]
    tok = d["text"][last].strip()
    # the wrecked token shows up as a bare number - or as a scrap of "Arabic" at the end of a Bangla line
    digit_tail = re.fullmatch(r"[0-9০-৯]{1,4}\W*", tok) is not None or bool(AR_RE.search(tok))
    short_end = d["left"][last] + d["width"][last] < 0.85 * img.size[0] and not tok.endswith(tuple("।?!”’\"»)"))
    if not (digit_tail or short_end):
        return
    if tess_cmd:
        pytesseract.pytesseract.tesseract_cmd = tess_cmd
    l, t, w, h = d["left"][last], d["top"][last], d["width"][last], d["height"][last]
    crop = ImageOps.expand(img.crop((l, max(0, t - 4), l + int(w * 0.8), t + h + 8)), border=20, fill=255)
    cfg = "--psm 7" + (f' --tessdata-dir "{tessdata}"' if tessdata else "")
    new = pytesseract.image_to_string(crop, lang="ben", config=cfg).strip()
    core = new.rstrip("।|!?”’\"*")
    if not BN_RE.search(new) or not core:
        return
    # a word that was read is only replaced when the re-read KEEPS it whole and adds the lost sentence end
    if digit_tail or (core.startswith(tok) and new != tok and new[-1] in "।|!?”’"):
        d["text"][last] = new


def _read_line(img, box, tess_cmd, tessdata):
    """OCR one line crop on its own with the Bangla model -> (text, mean confidence)."""
    import pytesseract
    from PIL import ImageOps
    if tess_cmd:
        pytesseract.pytesseract.tesseract_cmd = tess_cmd
    cfg = "--psm 7" + (f' --tessdata-dir "{tessdata}"' if tessdata else "")
    d = pytesseract.image_to_data(ImageOps.expand(img.crop(box), border=30, fill=255), lang="ben", config=cfg,
                                  output_type=pytesseract.Output.DICT)
    words = [(t, float(c)) for t, c in zip(d["text"], d["conf"]) if t.strip()]
    if not words:
        return "", 0.0
    return " ".join(t for t, _ in words), sum(c for _, c in words) / len(words)


def _second_look(img, lines, tess_cmd, tessdata):
    """Two repairs of the page-level OCR, both by reading a strip of the page as ONE line:
      * the short last line of a paragraph (where the footnote superscript sits) is read again; the new reading is
        kept when it is at least as confident and loses nothing ('দয়া করছে' -> 'দয়া করছো।');
      * a strip with ink that the page OCR skipped completely (a lone 'হাঁ।' under a paragraph) is read and added."""
    W, H = img.size
    if not lines:
        return
    med_h = statistics.median([l["bottom"] - l["top"] for l in lines])
    for ln in lines:
        txt = ln["text"]
        if AR_RE.search(txt) or len(BN_RE.findall(txt)) < 4 or ln["right"] >= 0.85 * W \
                or not 0.5 * med_h <= ln["bottom"] - ln["top"] <= 1.8 * med_h:
            continue
        new, conf = _read_line(img, (max(0, ln["left"] - 15), max(0, ln["top"] - 8), min(W, ln["right"] + 60),
                                     min(H, ln["bottom"] + 8)), tess_cmd, tessdata)
        if conf >= ln["conf"] - 2 and len(BN_RE.findall(new)) >= len(BN_RE.findall(txt)) and not AR_RE.search(new):
            ln["text"], ln["conf"] = new, max(conf, ln["conf"])
    covered = [False] * H
    for ln in lines:
        for y in range(max(0, ln["top"] - 4), min(H, ln["bottom"] + 4)):
            covered[y] = True
    grey = img.load()
    step = max(1, W // 600)                               # sample columns: fast enough in pure Python
    xs = range(0, W, step)
    bands, start = [], None
    for y in range(int(0.08 * H), int(0.92 * H)):
        dark = (not covered[y]) and sum(1 for x in xs if grey[x, y] < 128) > 2
        if dark and start is None:
            start = y
        elif not dark and start is not None:
            bands.append((start, y))
            start = None
    for t, b in bands:
        if not 0.4 * med_h <= b - t <= 1.6 * med_h:
            continue
        cols = [x for x in xs if any(grey[x, y] < 128 for y in range(t, b, 2))]
        if not cols or cols[-1] - cols[0] > 0.6 * W:      # a whole missed line is Arabic or a rule, not a stray word
            continue
        box = (max(0, cols[0] - 20), max(0, t - 8), min(W, cols[-1] + 20), min(H, b + 8))
        text, conf = _read_line(img, box, tess_cmd, tessdata)
        if conf >= 75 and len(BN_RE.findall(text)) >= 2 and not AR_RE.search(text):
            crop = img.crop(box)
            hist = crop.histogram()
            lines.append(dict(text=text, left=box[0], top=t, right=box[2], bottom=b, conf=conf,
                              ink=sum(hist[:128]) / max(1, sum(hist))))


def _ocr_worker(args):
    page, img_path, json_path, lang, tess_cmd, tessdata = args
    import pytesseract
    from PIL import Image
    if tess_cmd:
        pytesseract.pytesseract.tesseract_cmd = tess_cmd
    cfg = "--psm 4"
    if tessdata:
        cfg += f' --tessdata-dir "{tessdata}"'
    img = Image.open(img_path).convert("L")
    W, H = img.size
    d = pytesseract.image_to_data(img, lang=lang, config=cfg, output_type=pytesseract.Output.DICT)
    groups = defaultdict(list)
    for i, t in enumerate(d["text"]):
        if t.strip() and d["level"][i] == 5:
            groups[(d["block_num"][i], d["par_num"][i], d["line_num"][i])].append(i)
    for idx in groups.values():
        _fix_tail(img, d, idx, tess_cmd, tessdata)
    lines = []
    for idx in groups.values():
        l = min(d["left"][i] for i in idx)
        t = min(d["top"][i] for i in idx)
        r = max(d["left"][i] + d["width"][i] for i in idx)
        b = max(d["top"][i] + d["height"][i] for i in idx)
        confs = [float(d["conf"][i]) for i in idx if float(d["conf"][i]) >= 0]
        crop = img.crop((l, t, r, b))
        hist = crop.histogram()                      # 256 grey levels; "ink" = share of dark pixels
        ink = sum(hist[:128]) / max(1, sum(hist))
        lines.append(dict(text=" ".join(d["text"][i] for i in idx), left=l, top=t, right=r, bottom=b,
                          conf=sum(confs) / len(confs) if confs else 0.0, ink=ink))
    _second_look(img, lines, tess_cmd, tessdata)
    lines.sort(key=lambda x: x["top"])
    with open(json_path, "w", encoding="utf8") as f:
        json.dump(dict(page=page, width=W, height=H, image=os.path.basename(img_path), lines=lines), f,
                  ensure_ascii=False)
    return page


def run_ocr(paths, ocr_dir, args):
    from concurrent.futures import ProcessPoolExecutor, as_completed
    os.makedirs(ocr_dir, exist_ok=True)
    todo = []
    for p, ip in sorted(paths.items()):
        jp = os.path.join(ocr_dir, f"page_{p:03d}.json")
        if not os.path.exists(jp):
            todo.append((p, ip, jp, args.lang, args.tesseract_cmd, args.tessdata))
    print(f"OCR: {len(paths) - len(todo)} pages cached, {len(todo)} to do, {args.workers} workers")
    if not todo:
        return
    t0 = time.time()
    done = 0
    with ProcessPoolExecutor(max_workers=args.workers) as ex:
        futs = [ex.submit(_ocr_worker, t) for t in todo]
        for fu in as_completed(futs):
            fu.result()
            done += 1
            el = time.time() - t0
            eta = el / done * (len(todo) - done)
            print(f"  page done {done}/{len(todo)}  elapsed {el/60:.1f} min  ETA {eta/60:.1f} min", flush=True)


def find_rule(img_path):
    """y of the short horizontal rule printed above the footnotes (None if the page has none).
    Bangla headlines are long thin strokes too, so a candidate must also be ISOLATED: blank just above and below."""
    import cv2
    import numpy as np
    g = cv2.imread(img_path, cv2.IMREAD_GRAYSCALE)
    H, W = g.shape
    bw = (g < 128).astype(np.uint8)
    lines = cv2.morphologyEx(bw, cv2.MORPH_OPEN, np.ones((1, max(10, int(0.15 * W))), np.uint8))
    rows = np.where(lines.sum(axis=1) >= 0.15 * W)[0]
    rows = rows[(rows > 0.5 * H) & (rows < 0.97 * H)]
    if not len(rows):
        return None
    groups, cur = [], [rows[0]]
    for y in rows[1:]:
        if y - cur[-1] <= 1:
            cur.append(y)
        else:
            groups.append(cur)
            cur = [y]
    groups.append(cur)
    for gp in groups:
        top, bot = gp[0], gp[-1]
        if bot - top > 12:                                # too thick for a rule
            continue
        xs = np.where(lines[top:bot + 1].sum(axis=0) > 0)[0]
        if not 0.15 * W <= xs[-1] - xs[0] <= 0.7 * W:
            continue
        above = bw[max(0, top - 14):max(0, top - 3), xs[0]:xs[-1] + 1].sum()
        below = bw[bot + 4:bot + 15, xs[0]:xs[-1] + 1].sum()
        if above < 0.02 * (xs[-1] - xs[0]) * 11 and below < 0.02 * (xs[-1] - xs[0]) * 11:
            return int(top)
    return None


# =============================================================================== 3. classify
def _clearly_bangla(ln):
    """Readable Bangla with not one Arabic letter in it: never part of an Arabic block, however light it looks."""
    return not AR_RE.search(ln["text"]) and len(BN_RE.findall(ln["text"])) >= 6 and ln["conf"] >= 62


def classify_page(data, C):
    W, H = data["width"], data["height"]
    lines = data["lines"]
    for ln in lines:
        txt = ln["text"]
        a, b = len(AR_RE.findall(txt)), len(BN_RE.findall(txt))
        ln["ar_ratio"] = a / (a + b) if (a + b) else 0.0
        ln["h"] = ln["bottom"] - ln["top"]
        ln["cy"] = (ln["top"] + ln["bottom"]) / 2 / H
        ln["kind"] = None
    # reference statistics from clean, confident Bangla body lines (no Arabic characters at all)
    ref = [l for l in lines if l["cy"] >= C["header_zone"] and l["cy"] < C["footnote_zone"]
           and not AR_RE.search(l["text"]) and l["conf"] >= 75 and len(l["text"]) >= 25]
    med_h = statistics.median([l["h"] for l in ref]) if ref else 0
    med_ink = statistics.median([l["ink"] for l in ref]) if ref else 0
    # header / noise / arabic
    for ln in lines:
        txt = ln["text"].strip()
        n_ar = len(AR_RE.findall(txt))
        light = med_ink and ln["ink"] < C["arabic_ink_ratio"] * med_ink      # Arabic strokes are lighter
        if ln["cy"] < C["header_zone"]:
            ln["kind"] = "header"
        elif ln["cy"] > C["ornament_zone"] and (len(txt) <= 4 or ln["conf"] < 40):
            ln["kind"] = "noise"
        elif ln["ar_ratio"] >= C["arabic_ratio"] and n_ar >= 3:
            ln["kind"] = "arabic"
        elif light and ln["ar_ratio"] >= C["arabic_ratio_if_light"] and n_ar >= 3:
            ln["kind"] = "arabic"
        elif len(txt) <= 3 and ln["cy"] < C["footnote_zone"]:
            ln["kind"] = "noise"                                              # stray specks / superscripts
        elif ln["conf"] < 15 and len(txt) < 4:
            ln["kind"] = "noise"
        elif med_h and ln["h"] < 0.25 * med_h and len(txt) <= 3:
            ln["kind"] = "noise"                                              # 1-2 px specks anywhere on the page
    # smoothing: a line between two Arabic lines is Arabic too (translation never sits inside an Arabic block)
    seq = [l for l in lines if l["kind"] not in ("header", "noise")]
    for i in range(1, len(seq)):
        cur, prv = seq[i], seq[i - 1]
        if cur["kind"] is not None or prv["kind"] != "arabic":
            continue
        nxt_ar = i + 1 < len(seq) and seq[i + 1]["kind"] == "arabic"
        light = med_ink and cur["ink"] < C["arabic_ink_ratio"] * med_ink
        if _clearly_bangla(cur):
            continue
        # between two Arabic lines, or the short last line of an Arabic block (light strokes, shaky OCR)
        if nxt_ar or (light and cur["conf"] < 70 and not MARKER_RE.match(cur["text"].strip())):
            cur["kind"] = "arabic"
    for i in range(len(seq) - 1):                  # first line of an Arabic block misread as Bangla
        cur = seq[i]
        if cur["kind"] is None and seq[i + 1]["kind"] == "arabic" and not MARKER_RE.match(cur["text"].strip()):
            light = med_ink and cur["ink"] < C["arabic_ink_ratio"] * med_ink
            if light and cur["conf"] < 70 and not _clearly_bangla(cur):
                cur["kind"] = "arabic"
    # footnotes: everything below the printed footnote rule; without a rule only unmistakable footnote lines count
    start = None
    rule = data.get("rule_y")
    for i, ln in enumerate(lines):
        if ln["kind"] is not None:
            continue
        txt = ln["text"].strip()
        if rule is not None:
            if ln["top"] >= rule - 5:
                start = i
                break
        elif ln["cy"] > C["footnote_zone"] and not MARKER_RE.match(txt):
            tiny_number = ln["h"] < 0.6 * med_h and re.fullmatch(r"\W*[০-৯0-9]{1,4}\W*", txt)
            foot_like = re.match(r"^\W{0,4}[০-৯0-9]{0,4}\s*[.।]\s+\S", txt) and "হাদিসের মান" in txt
            if tiny_number or foot_like or txt.startswith("হাদিসের মান"):
                start = i
                break
    if start is not None:
        for ln in lines[start:]:
            if ln["kind"] is None:
                ln["kind"] = "footnote"
    # headings (bold = much more ink per pixel) vs body
    for ln in lines:
        if ln["kind"] is not None:
            continue
        txt = ln["text"].strip()
        bold = med_ink and ln["ink"] >= C["heading_ink_ratio"] * med_ink
        huge = med_h and ln["h"] >= C["heading_huge_ratio"] * med_h and ln["ink"] >= 1.1 * med_ink   # tall AND not lighter than text
        cx = (ln["left"] + ln["right"]) / 2 / W
        wide = (ln["right"] - ln["left"]) / W >= 0.6
        centered = abs(cx - 0.5) < C["heading_center_tol"]      # headings are centred; short last body lines are not
        ln["ink_r"] = ln["ink"] / med_ink if med_ink else 0.0
        ln["h_r"] = ln["h"] / med_h if med_h else 0.0
        if (bold or huge) and (wide or centered) and len(txt) <= C["heading_max_chars"] \
                and not MARKER_RE.match(txt):
            ln["kind"] = "heading"
        else:
            ln["kind"] = "body"
    return lines


# =============================================================================== 4. segment
def clean_bn(text, strip_lead=False):
    """Remove OCR junk from Bangla text: stray Arabic letters, direction marks, short Latin specks."""
    text = AR_RE.sub("", text)
    text = re.sub(r"[‎‏‪-‮]", "", text)
    text = re.sub(r"(?<=[ঀ-৿])\s*[|Il](?=\s|$|[”’\"\)\]]|[০-৯0-9*])", "।", text)   # the danda is often read as a pipe
    text = re.sub(r"(?<=[।?!\"”’»])\s*[০-৯]{2,4}(?=\s|$)", "", text)   # footnote superscripts after sentence end
    text = re.sub(r"(?:(?<=\s)|^)[A-Za-z0-9°*&@#~|\\/_^]{1,3}(?=\s|$)", " ", text)  # short Latin/symbol specks
    text = re.sub(r"\*+[০-৯0-9]{0,3}", "", text)                         # footnote reference marks: "*" / "*২"
    text = re.sub(r"[-_–—]*_[-_–—]*|--+", "—", text)                       # the print's dash "—" is read as "-_" / "--" / "_"
    text = common_fixes(text)
    text = re.sub(r"\s+", " ", text).strip()
    if strip_lead:
        text = re.sub(r"^[^ঀ-৿]+", "", text)
    return text


def _closes_reference(text):
    """True when the text ends with the number that closes an open bracket: "(সুরা মুমতাহিনা : ৮)"."""
    m = re.search(r"[০-৯]{1,3}(?:\s*[-,]\s*[০-৯]{1,3})?\s*\)\W{0,2}$", text)
    return bool(m) and text[:m.start()].count("(") > text[:m.start()].count(")")


def trim_tail(text):
    """Drop OCR junk after the last sentence end ('।” ৬9 চল চট 0৪৪ উস 2০41'): a short tail holding Latin letters or
    digit mixtures cannot be Bangla text."""
    # a trailing footnote number - but not the number that closes a reference such as "(সুরা মুমতাহিনা : ৮)"
    if re.search(r"[ঀ-৿]\W{0,3}\s*[০-৯0-9]{1,4}\W{0,2}$", text) and not _closes_reference(text):
        text = re.sub(r"\s*[০-৯0-9]{1,4}\W{0,2}$", "", text)
    text = re.sub(r"(?<=[।?!”’)])['\"]+$", "", text)             # stray quote from the superscript
    while True:                                           # trailing tokens with digits/Latin and no Bangla letter
        m = re.search(r"\s+(\S+)$", text)
        if m and not re.search(r"[\u0985-\u09B9\u09DC-\u09DF]", m.group(1)) \
                and re.search(r"[0-9০-৯A-Za-z<>]", m.group(1)) and not _closes_reference(text):
            text = text[:m.start()]
        else:
            break
    m = None
    for m in re.finditer(r"[।?!”’\"»)]", text):
        pass
    if m and 0 < len(text) - m.end() <= 45:
        tail = text[m.end():]
        if re.search(r"[A-Za-z]|[০-৯][0-9]|[0-9][০-৯]|\d{3,}", tail):
            return text[:m.end()].rstrip()
    return text


def clean_ar(text):
    """Remove Bangla/Latin junk that OCR sometimes produces inside Arabic lines."""
    text = re.sub(r"[ঀ-৿A-Za-z]", "", text)
    text = re.sub(r"[‎‏‪-‮]", "", text)
    return re.sub(r"\s+", " ", text).strip()


def build_entries(page_data, C):
    entries, unnumbered, foot = [], [], defaultdict(list)
    cur = None
    pending_ar, pending_ar_pages, pending_conf = [], [], []
    heading, prev_heading_line, chapter = "", False, ""
    fblocks, foot_open = [], False         # footnote blocks in reading order

    def new_entry(no, page, flag=None):
        e = dict(no=no, heading=heading, chapter=chapter, arabic=[], bangla=[], pages=[page], conf=[], flags=[])
        if flag:
            e["flags"].append(flag)
        if pending_ar:
            e["arabic"] = list(pending_ar)
            e["pages"] = sorted(set(pending_ar_pages + [page]))
            e["conf"] += pending_conf
        entries.append(e)
        return e

    for pd in page_data:
        pg = pd["page"]
        for ln in pd["lines"]:
            k, txt = ln["kind"], ln["text"].strip()
            if k in ("header", "noise"):
                continue
            if k == "footnote":
                frag = re.fullmatch(r"\W{0,3}([০-৯0-9]{1,4})\W{0,3}", txt)
                if frag:                                   # tiny superscript = start of a new footnote
                    fblocks.append(dict(no=bn_int(frag.group(1)), page=pg, parts=[]))
                    foot_open = True
                    continue
                m = FOOT_START_RE.match(txt)
                if m:
                    fblocks.append(dict(no=bn_int(m.group(1)), page=pg, parts=[txt[m.end():]]))
                    foot_open = True
                elif foot_open and fblocks:
                    fblocks[-1]["parts"].append(txt)
                else:                                      # superscript was merged into the text: number lost
                    fblocks.append(dict(no=None, page=pg, parts=[txt]))
                    foot_open = True
                if re.search(r"হাদিসের মান\s*[:ঃ]?\s*[ঀ-৿]", " ".join(fblocks[-1]["parts"])):
                    foot_open = False                      # grade word present = footnote complete
                    # (if the grade word wrapped onto the next line, the block stays open until it arrives)
                continue
            foot_open = False
            if k == "chapter":                         # "অধ্যায় : ..." printed in the text itself
                chapter, prev_heading_line = ln.get("chapter_name") or txt, False
                continue
            if k == "heading":
                heading = (heading + " " + txt) if prev_heading_line else txt
                prev_heading_line = True
                continue
            prev_heading_line = False
            if k == "arabic":
                # Arabic after a translation is either the NEXT entry's Arabic (a [number] marker follows)
                # or a quote inside the current translation (more Bangla follows) - decided at the next body line.
                pending_ar.append(txt)
                pending_ar_pages.append(pg)
                pending_conf.append(ln["conf"])
                continue
            # body line
            m = MARKER_RE.match(txt)
            if not m and pending_ar and (len(txt) <= 4 or (ln["conf"] < 45 and len(txt) < 25)):
                continue                            # OCR speck after an Arabic block, not real translation text
            if m:
                cur = new_entry(bn_int(m.group(1) or m.group(2)), pg)
                cur["anchor"] = dict(page=pg, y0=round(ln["top"] / pd["height"], 4),
                                     y1=round(ln["bottom"] / pd["height"], 4))   # where the hadith starts (web review)
                pending_ar, pending_ar_pages, pending_conf = [], [], []
                txt = txt[m.end():]
            elif cur is None:
                if pending_ar and (len(txt) < 12 or ln["conf"] < 50):
                    continue                        # speck / junk between Arabic and its marker
                if pending_ar:                      # Arabic followed by Bangla with no marker
                    cur = new_entry(None, pg, "no_marker")
                    cur["anchor"] = dict(page=pg, y0=round(ln["top"] / pd["height"], 4),
                                         y1=round(ln["bottom"] / pd["height"], 4))
                    pending_ar, pending_ar_pages, pending_conf = [], [], []
                else:                               # front matter / intro text
                    unnumbered.append((pg, "text", txt))
                    continue
            elif pending_ar:                        # Arabic quote inside the current translation
                cur["arabic"] += pending_ar
                pending_ar, pending_ar_pages, pending_conf = [], [], []
            cur["bangla"].append(txt)
            cur["conf"].append(ln["conf"])
            if pg not in cur["pages"]:
                cur["pages"].append(pg)

    # footnotes: drop junk blocks, fill lost numbers (footnote numbers run consecutively), index by number
    prev_no = None
    for b in fblocks:
        b["text"] = clean_bn(" ".join(b["parts"]), strip_lead=True)
        if len(BN_RE.findall(b["text"])) < 6:
            b["junk"] = True
            continue
        if b["no"] is None and prev_no is not None:
            b["no"] = prev_no + 1
            b["guessed"] = True
        if b["no"] is not None:
            prev_no = b["no"]
    good = [b for b in fblocks if not b.get("junk")]
    nxt_no = None
    for b in reversed(good):                               # leading blocks with no predecessor: count back
        if b["no"] is None and nxt_no is not None and nxt_no > 1:
            b["no"], b["guessed"] = nxt_no - 1, True
        if b["no"] is not None:
            nxt_no = b["no"]
    for b in good:
        foot[b["no"]].append(b)

    # finalise
    prev = None
    for e in entries:
        e["arabic_text"] = clean_ar(" ".join(e["arabic"]))
        e["bangla_text"] = trim_tail(clean_bn(" ".join(e["bangla"])))
        e["conf_mean"] = round(sum(e["conf"]) / len(e["conf"]), 1) if e["conf"] else 0
        e["page_start"], e["page_end"] = min(e["pages"]), max(e["pages"])
        e["foot_text"] = ""
        if e["no"] is not None and e["no"] in foot:
            blocks = foot.pop(e["no"])
            e["foot_text"] = " ".join(b["text"] for b in blocks)
        if e["no"] is not None:
            if prev is not None and e["no"] != prev + 1:
                e["flags"].append(f"seq_gap(expected {prev + 1})")
            prev = e["no"]
        if e["conf_mean"] < C["low_conf"]:
            e["flags"].append("low_conf")
        if len(e["bangla_text"]) < C["min_text_chars"]:
            e["flags"].append("short_text")
        if not e["arabic_text"]:
            e["flags"].append("no_arabic")
    for no, blocks in foot.items():
        for b in blocks:
            unnumbered.append((b["page"], f"footnote {no if no is not None else '?'} (unmatched)", b["text"]))
    return entries, unnumbered


def _norm(t):
    """Bangla letters + vowel signs only (re's \\W would also strip the vowel signs)."""
    return re.sub(r"[০-৯]+", "", re.sub(r"[^\u0980-\u09FF]+", "", t))


def apply_toc_headings(page_data, titles, C):
    """Body-page headings are only slightly bolder than the text, so the ink test misses some and they leak into
    the previous hadith. A bold-ish body line that matches a TOC section title (or a wrapped piece of one) IS a heading."""
    from difflib import SequenceMatcher
    tn = [_norm(t) for t in titles if len(_norm(t)) >= 5]
    n = 0
    for pd in page_data:
        for ln in pd["lines"]:
            if ln["kind"] != "body" or not (ln.get("ink_r", 0) >= C["toc_head_ink"] or ln.get("h_r", 0) >= C["toc_head_h"]):
                continue
            txt = ln["text"].strip()
            if MARKER_RE.match(txt):
                continue
            ln_n = _norm(txt)
            if len(ln_n) < 5:
                continue
            for t in tn:
                sm = SequenceMatcher(None, ln_n, t, autojunk=False)
                if sm.quick_ratio() < 0.6:
                    continue
                m = sm.find_longest_match(0, len(ln_n), 0, len(t))
                if m.size >= 0.75 * len(ln_n) and m.size >= 5 or sm.ratio() >= 0.8:
                    ln["kind"] = "heading"
                    n += 1
                    break
    return n


CHAPTER_GARBLE_RE = re.compile(r"^\W*[0-9]{4,}\W*[:：]")    # the decorative "অধ্যায়" is often read as digits


def structural_headings(page_data, C):
    """Headings the ink test missed, found by LAYOUT: a centred line that is indented on both sides, a little bolder
    than the text, and directly followed by the Arabic of the next hadith (or by another heading line)."""
    n = 0
    for pd in page_data:
        W = pd["width"]
        wide = [l for l in pd["lines"] if l["kind"] == "body" and (l["right"] - l["left"]) / W > 0.6]
        if len(wide) < 3:
            continue
        ml = statistics.median([l["left"] for l in wide])
        mr = statistics.median([l["right"] for l in wide])
        seq = [l for l in pd["lines"] if l["kind"] not in ("header", "noise")]
        for i in range(len(seq) - 2, -1, -1):
            ln, nxt = seq[i], seq[i + 1]
            txt = ln["text"].strip()
            if ln["kind"] != "body" or nxt["kind"] not in ("arabic", "heading") or MARKER_RE.match(txt):
                continue
            cx = (ln["left"] + ln["right"]) / 2
            if ln.get("ink_r", 0) >= C["struct_head_ink"] and ln["left"] >= ml + 0.02 * W \
                    and ln["right"] <= mr - 0.02 * W and abs(cx - (ml + mr) / 2) <= 0.04 * W \
                    and 4 <= len(txt) <= C["heading_max_chars"]:
                ln["kind"] = "heading"
                n += 1
        # the short last line of a wrapped heading, taken for Arabic only because it is light and badly read
        for i in range(1, len(seq)):
            ln = seq[i]
            cx = (ln["left"] + ln["right"]) / 2
            if ln["kind"] == "arabic" and seq[i - 1]["kind"] == "heading" and not AR_RE.search(ln["text"]) \
                    and len(BN_RE.findall(ln["text"])) >= 4 and ln["left"] >= ml + 0.02 * W \
                    and ln["right"] <= mr - 0.02 * W and abs(cx - (ml + mr) / 2) <= 0.04 * W:
                ln["kind"] = "heading"
                n += 1
    return n


def reread_headings(page_data, img_dir, a):
    """Section names when there is no table of contents to copy them from: every heading line is read several
    more times on its own (Bangla model, different sizes) and the reading whose words actually occur in the book's
    running text wins - the running text is read far more reliably than the bold headings."""
    import cv2
    import numpy as np
    import pytesseract
    from collections import Counter
    from PIL import Image, ImageOps
    if a.tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = a.tesseract_cmd
    cfg = "--psm 7" + (f' --tessdata-dir "{a.tessdata}"' if a.tessdata else "")
    word_re = re.compile(r"[\u0980-\u09FF]{2,}")
    vocab = {w for pd in page_data for l in pd["lines"] if l["kind"] == "body" and l["conf"] >= 80
             for w in word_re.findall(l["text"])}

    def score(t):
        words = word_re.findall(t)
        return sum(w in vocab for w in words) / len(words) if words else 0.0

    for pd in page_data:
        heads = [l for l in pd["lines"] if l["kind"] == "heading"]
        if not heads:
            continue
        g = cv2.imread(os.path.join(img_dir, pd["image"]), cv2.IMREAD_GRAYSCALE)
        H, W = g.shape
        for ln in heads:
            x0, x1 = max(0, ln["left"] - 25), min(W, ln["right"] + 25)
            y0, y1 = max(0, ln["top"] - 10), min(H, ln["bottom"] + 10)
            rows = (g[y0:y1, x0:x1] < 128).sum(axis=1) > 0
            runs, start, gap = [], None, 0                # the line's box may catch the top of the next line:
            for i, on in enumerate(list(rows) + [False] * 6):   # keep the tallest band of ink only
                if on:
                    start, gap = (i if start is None else start), 0
                elif start is not None:
                    gap += 1
                    if gap >= 5:
                        runs.append((start, i - gap + 1))
                        start = None
            if runs:
                t, b = max(runs, key=lambda r: r[1] - r[0])
                y0, y1 = max(0, y0 + t - 8), min(H, y0 + b + 8)
            band = g[y0:y1, x0:x1]
            cands = [ln["text"].strip()]
            for sc in (1.0, 0.85, 0.75, 0.6):
                im = band if sc == 1 else cv2.resize(band, None, fx=sc, fy=sc, interpolation=cv2.INTER_AREA)
                t = pytesseract.image_to_string(ImageOps.expand(Image.fromarray(im), border=30, fill=255),
                                                lang="ben", config=cfg).strip()
                if t and not AR_RE.search(t):
                    cands.append(t)
            longest = max(len(BN_RE.findall(c)) for c in cands)
            cands = [c for c in cands if len(BN_RE.findall(c)) >= 0.7 * longest]     # no clipped readings
            votes = Counter(cands)
            ln["text"] = max(cands, key=lambda c: (round(score(c), 3), votes[c], c == cands[0]))


def mark_chapters(page_data, img_dir, a, known, reread):
    """Heading lines that are a chapter title ("অধ্যায় : ...") become kind 'chapter'. With reread=True (no table of
    contents to take the name from) the line is OCRed again on its own and matched against the known chapter names."""
    from difflib import SequenceMatcher
    names = []
    for pd in page_data:
        for ln in pd["lines"]:
            txt = ln["text"].strip()
            if not (CHAPTER_RE.match(txt) or CHAPTER_GARBLE_RE.match(txt)):
                continue
            if not (ln["kind"] == "heading" or (ln["kind"] == "body" and ln.get("ink_r", 0) >= 1.25)):
                continue
            ln["kind"] = "chapter"
            if not reread:
                continue
            import cv2
            import pytesseract
            from PIL import Image, ImageOps
            if a.tesseract_cmd:
                pytesseract.pytesseract.tesseract_cmd = a.tesseract_cmd
            tdir = f' --tessdata-dir "{a.tessdata}"' if a.tessdata else ""
            g = cv2.imread(os.path.join(img_dir, pd["image"]), cv2.IMREAD_GRAYSCALE)
            band = g[max(0, ln["top"] - 12):ln["bottom"] + 12]
            cands = []
            for scale in (0.6, 1.0):                           # the decorative type reads best a bit smaller
                im = cv2.resize(band, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA) if scale != 1 else band
                im = ImageOps.expand(Image.fromarray(im), border=30, fill=255)
                for psm in (13, 7):
                    t = pytesseract.image_to_string(im, lang="ben", config=f"--psm {psm}" + tdir).strip()
                    t = clean_toc_name(re.sub(r"^[^:：ঃ]{0,14}[:：ঃ]\s*", "", t))
                    if len(BN_RE.findall(t)) >= 4:
                        cands.append(t)
            best = (0.0, None)
            for t in cands:
                for k in known:
                    best = max(best, (SequenceMatcher(None, _norm(t), _norm(k), autojunk=False).ratio(), k))
            if best[0] >= 0.5:
                ln["chapter_name"] = best[1]
            else:
                ln["chapter_name"] = max(cands, key=len) if cands else "অধ্যায়"
                ln["chapter_unsure"] = True
            names.append(ln["chapter_name"])
    return names


def structure_from_headings(entries):
    """No table of contents: chapters and sections are the headings met in the text, in order.
    -> chapters, sections, chap_of (per hadith, -1 = none), sec_of (per hadith, -1 = none)."""
    chapters, sections, chap_of, sec_of = [], [], [], []
    prev = None
    for e in entries:
        key = (e.get("chapter", ""), e["heading"])
        if key != prev:
            if e.get("chapter") and (prev is None or prev[0] != e["chapter"]):
                chapters.append(e["chapter"])
            name = clean_toc_name(e["heading"])
            if name:
                sections.append(name)
            prev = key
            cur_sec = len(sections) - 1 if name else -1
        chap_of.append(len(chapters) - 1)
        sec_of.append(cur_sec)
    return chapters, sections, chap_of, sec_of


def running_title(page_data):
    """The book title as printed in the running header of the pages (None if there is no consistent header)."""
    from collections import Counter
    c = Counter(re.sub(r"\s+", " ", l["text"]).strip() for pd in page_data for l in pd["lines"]
                if l["kind"] == "header" and len(BN_RE.findall(l["text"])) >= 4)
    if c:
        text, n = c.most_common(1)[0]
        if n >= max(2, len(page_data) // 4):
            return text
    return None


def detect_toc_pages(pdf, img_dir, out_dir, n_pages):
    """PDF pages that are a table of contents = pages full of dotted leaders ("title ........ 27"). [] if none."""
    cache = os.path.join(out_dir, "toc_pages.json")
    if os.path.exists(cache):
        return json.load(open(cache))
    first = list(range(1, min(n_pages, 60) + 1))
    paths = extract_images(pdf, img_dir, first)
    hits = []
    for p in first:
        ys = sorted(y for y, _, cnt in _strip_leaders(paths[p])[1] if cnt >= 8)
        rows = sum(1 for i, y in enumerate(ys) if i == 0 or y - ys[i - 1] > 15)
        if rows >= 8:
            hits.append(p)
    best, run = [], []
    for p in hits:                                            # longest run of consecutive pages
        run = run + [p] if run and p == run[-1] + 1 else [p]
        if len(run) > len(best):
            best = run
    json.dump(best, open(cache, "w"))
    return best


def _same_number(a, b):
    """Equal, or equal up to the classic Bangla OCR confusion of ৪ (4) and ৮ (8)."""
    sa, sb = str(a), str(b)
    return len(sa) == len(sb) and all(x == y or {x, y} == {"4", "8"} for x, y in zip(sa, sb))


def repair_numbers(entries):
    """Hadith numbers are 1,2,3,... so the previous number predicts the next one. Fix what OCR got wrong:
      * ৪/৮ mix-ups ([৮০০] for [৪০০]) are corrected silently;
      * a marker that does not fit but whose NEXT neighbour proves it is the expected number is corrected
        (flag no_repaired);
      * a marker that jumps forward a little is kept and flagged seq_gap (a hadith in between was not detected);
        a huge jump is a misread and gets the expected number (flag no_repaired);
      * a marker that jumps backward is not a hadith start ("[৭]" inside a list...): its text is appended to
        the previous hadith (flag merged_marker)."""
    out, prev = [], None
    for i, e in enumerate(entries):
        e["flags"] = [f for f in e["flags"] if not f.startswith("seq_gap")]
        o = e["no"]
        if prev is None:
            out.append(e)
            prev = _split_embedded(e, out)["no"]
            continue
        exp = prev + 1
        nxt = entries[i + 1]["no"] if i + 1 < len(entries) else None
        if o == exp:
            pass
        elif o is not None and _same_number(o, exp):
            e["no"] = exp
        elif o is None or o > exp + C_MAX_JUMP or (nxt is not None and (nxt == exp + 1 or _same_number(nxt, exp + 1))):
            e["flags"].append(f"no_repaired(read {o})")
            e["no"] = exp
        elif o > exp:
            e["flags"].append(f"seq_gap(expected {exp})")
        else:                                                 # backwards: not a new hadith
            if len(e["bangla_text"]) >= C_MIN_KEEP:
                out[-1]["bangla_text"] = (out[-1]["bangla_text"] + " " + e["bangla_text"]).strip()
                out[-1]["page_end"] = max(out[-1]["page_end"], e["page_end"])
                out[-1]["flags"].append(f"merged_marker(read {o})")
            continue
        out.append(e)
        e = _split_embedded(e, out)
        prev = e["no"]
    entries[:] = out


def _split_embedded(e, out):
    """The previous hadith ended in the middle of a printed line, so the next marker sits mid-line: cut there."""
    while e["no"] is not None:
        m = EMBED_RE.search(e["bangla_text"], 10)
        n = bn_int(m.group(1)) if m else None
        if n is None or not (n == e["no"] + 1 or _same_number(n, e["no"] + 1)):
            break
        new = dict(e, no=e["no"] + 1, bangla_text=e["bangla_text"][m.end():].strip(),
                   page_start=e["page_end"], flags=["split_marker"], anchor=None)
        e["bangla_text"] = e["bangla_text"][:m.start()].strip()
        out.append(new)
        e = new
    return e


C_MAX_JUMP = 8             # forward jumps larger than this are misread digits, not missing hadiths
C_MIN_KEEP = 15            # a stray backwards marker shorter than this is junk and is dropped


# =============================================================================== 4b. table of contents
CHAPTER_RE = re.compile(r"^\s*অধ্যা(?:\u09df|\u09af\u09bc?)?\s*[:：ঃ.]?\s*")
TAIL_NUM_RE = re.compile(r"[\s.…,·:\-_*]*[০-৯0-9]{1,4}[\s.]*$")
TAIL_DOT_RE = re.compile(r"[\s.…,·:_*]+$")


def _strip_leaders(path):
    """Grey image with the dotted leaders ("......") of a TOC page erased + the (y, left x, dot count) of each leader.
    Leaders confuse Tesseract (it garbles the words next to them). Only EVENLY SPACED runs of >=5 small dots are
    removed, so the dot inside Bangla letters such as 'র' is kept."""
    import cv2
    import numpy as np
    g = cv2.imread(path, cv2.IMREAD_GRAYSCALE)
    _, bw = cv2.threshold(g, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    n, lab, st, cen = cv2.connectedComponentsWithStats(bw, connectivity=8)
    dots = sorted((i for i in range(1, n) if 3 <= st[i, 2] <= 16 and 3 <= st[i, 3] <= 16
                   and 0.5 < st[i, 2] / st[i, 3] < 2.0), key=lambda i: cen[i][1])
    groups, cur = [], []
    for i in dots:                                        # same baseline (+-3 px)
        if cur and abs(cen[i][1] - np.mean([cen[j][1] for j in cur])) > 3:
            groups.append(cur)
            cur = []
        cur.append(i)
    if cur:
        groups.append(cur)
    remove, leader_y = np.zeros(n, bool), []
    for gp in groups:
        gp.sort(key=lambda i: cen[i][0])
        runs, run = [], [gp[0]]
        for i in gp[1:]:
            if cen[i][0] - cen[run[-1]][0] <= 20:
                run.append(i)
            else:
                runs.append(run)
                run = [i]
        runs.append(run)
        for r in runs:
            if len(r) >= 5:
                leader_y.append((float(np.mean([cen[i][1] for i in r])), float(min(cen[i][0] for i in r)), len(r)))
                remove[r] = True
    mask = cv2.dilate(remove[lab].astype(np.uint8), np.ones((7, 7), np.uint8)) > 0     # also wipe the faint halo
    g[mask] = 255
    return g, leader_y


def _toc_rows(img_path, lang, tess_cmd, tessdata):
    """One dict per printed TOC line: text, vertical position, and whether a dotted leader was on that line.
    Lines WITH a leader are cut out around the leader's baseline and OCRed as a single line
    (--psm 7, so no layout analysis; the dots were erased beforehand). Lines without a leader (the first line of a
    wrapped title) are found by ink projection. Whole-page OCR kept dropping or merging bold chapter lines."""
    import pytesseract
    from PIL import Image, ImageOps
    if tess_cmd:
        pytesseract.pytesseract.tesseract_cmd = tess_cmd
    cfg = "--psm 7" + (f' --tessdata-dir "{tessdata}"' if tessdata else "")
    g, leaders = _strip_leaders(img_path)
    H, W = g.shape
    sc = W / 1781.0                                       # line geometry below was measured on a 1781 px wide scan

    def ocr(top, bot, x_end):
        crop = ImageOps.expand(Image.fromarray(g[max(0, top):min(H, bot), :max(40, x_end)]), border=30, fill=255)
        text = pytesseract.image_to_string(crop, lang=lang, config=cfg).strip()
        if not text:                                      # bold lines sometimes need to be a bit larger
            text = pytesseract.image_to_string(crop.resize((int(crop.width * 1.5), int(crop.height * 1.5))),
                                               lang=lang, config=cfg).strip()
        return text

    rows, wins = [], []
    clusters = []                                         # one leader per printed line: the longest dot run
    for y, x, n in sorted(leaders):
        if clusters and y - clusters[-1][0] <= 30 * sc:   # same printed line (lines are ~70 px apart)
            if n > clusters[-1][2]:
                clusters[-1] = (y, x, n)
        else:
            clusters.append((y, x, n))
    for y, x, n in clusters:
        top, bot = int(y - 50 * sc), int(y + 18 * sc)
        wins.append((top, bot))
        text = ocr(top, bot, W)
        # nothing readable on a leader line = a bold chapter line (they are the ones Tesseract gives up on);
        # keep a placeholder so the chapter COUNT and the sections under it stay right (name comes from chapters.txt)
        rows.append(dict(top=top, cy=(top + bot) / 2 / H, text=text or "অধ্যায় : ?", leader=True))
    rest = g.copy()                                       # lines without a leader: erase the ones done, find the rest
    for top, bot in wins:
        rest[max(0, top):bot] = 255
    ink = (rest < 128).sum(axis=1) > 3
    bands, start, gap = [], None, 0
    for y in range(H):
        if ink[y]:
            if start is None:
                start = y
            gap = 0
        elif start is not None:
            gap += 1
            if gap > 8:
                bands.append((start, y - gap))
                start = None
    for t, bt in bands:
        if bt - t < 25 * sc:                              # specks
            continue
        text = ocr(t - 8, bt + 9, W)
        if text:
            rows.append(dict(top=t, cy=(t + bt) / 2 / H, text=text, leader=False))
    rows.sort(key=lambda r: r["top"])
    return rows


TOC_FIXES = [                                         # systematic OCR slips seen in this book's TOC
    (r"যেব্যক্তি", "যে ব্যক্তি"), (r"বাক্তি", "ব্যক্তি"),
    (r"সম্[তভ]ান", "সন্তান"), (r"সাম্[তভ]না", "সান্ত্বনা"),
    (r"\sনে$", ""),                                       # a misread page number glued to a title
    # headings only (in running text these could be real words): the bold/small type drops the vowel sign ু
    (r"আয় (বৃদ্ধি|বাড়ে)", r"আয়ু \1"), (r"মাজলমের দআ", "মাজলুমের দুআ"), (r"জলম অন্ধকার", "জুলুম অন্ধকার"),
    (r"অসস্থকালেও সুস্থালের", "অসুস্থকালেও সুস্থকালের"), (r"গ্রামা ব্যক্তি", "গ্রাম্য ব্যক্তি"),
    (r"সমন্ভানাদি", "সন্তানাদি"), (r"ধনাচঢ্যতা", "ধনাঢ্যতা"), (r"সম্ভাব সৃষ্টি", "সদ্ভাব সৃষ্টি"),
    (r"বালিক'কে চুম দেয়া", "বালিকাকে চুমু দেয়া"), (r"নাবসা$", "না বসা"),
    (r"—\W*হে মুনাফিক", "—“হে মুনাফিক"),
]


def clean_toc_name(t):
    t = t.replace("\u200c", "").replace("\u200d", "")
    t = re.sub(r"[-_–—]*_[-_–—]*|--+", "—", t)            # "-_" "__" "--" were a dash in the print
    t = re.sub(r"(?<=[ঀ-৿])_(?=[ঀ-৿])", "-", t)           # "আচার_ব্যবহার" -> "আচার-ব্যবহার"
    for pat, rep in TOC_FIXES:
        t = re.sub(pat, rep, t)
    t = common_fixes(t)
    t = re.sub(r"^[^ঀ-৿\"“‘(]+", "", t)                  # leading specks
    t = re.sub(r"[^ঀ-৿?!।\"”’)]+$", "", t)               # trailing specks / leader remains
    return re.sub(r"\s+", " ", t).strip()


def read_toc(img_dir, pages, out_dir, a):
    """-> (chapters, sections, section_chapter, warnings); ids are assigned in reading order.
    section_chapter[i] = index of the chapter that section i is printed under."""
    cache = os.path.join(out_dir, "toc.json")
    if os.path.exists(cache):
        rows_by_page = {int(k): v for k, v in json.load(open(cache, encoding="utf8")).items()}
    else:
        rows_by_page = {}
    todo = [p for p in pages if p not in rows_by_page]
    if todo:
        paths = extract_images(a.pdf, img_dir, todo)
        for p in todo:
            rows_by_page[p] = _toc_rows(paths[p], "ben", a.tesseract_cmd, a.tessdata)
        json.dump(rows_by_page, open(cache, "w", encoding="utf8"), ensure_ascii=False)
    chapters, sections, sec_chapter, warns, parts = [], [], [], [], []
    for p in pages:
        for r in rows_by_page[p]:
            txt = r["text"].strip()
            if r["cy"] < C_TOC_HEADER or "চিপত্র" in txt or "আদাবুল" in txt:      # page header / "সূচিপত্র" title
                continue
            ends = r["leader"] or bool(TAIL_NUM_RE.search(txt))     # a leader or a trailing page number ends an entry
            body = TAIL_DOT_RE.sub("", TAIL_NUM_RE.sub("", txt)).strip() if ends else txt
            if CHAPTER_RE.match(txt):
                if parts:
                    warns.append(f"TOC page {p}: dropped unterminated text {' '.join(parts)!r}")
                    parts = []
                chapters.append(clean_toc_name(TAIL_DOT_RE.sub("", CHAPTER_RE.sub("", body))))
                continue
            if body:
                parts.append(body)
            if ends and parts:
                full = " ".join(parts)
                # a page number in the MIDDLE means two entries ran together (a wrapped title + the next one)
                for k, piece in enumerate(re.split(r"\s+[০-৯]{2,4}\s+", full)):
                    if k:
                        piece = re.sub(r"^\S{1,2}\s+", "", piece)      # speck in front of the second title
                    name = clean_toc_name(piece)
                    if len(BN_RE.findall(name)) < 4:
                        warns.append(f"TOC page {p}: dropped junk {name!r}")
                    else:
                        sections.append(name)
                        sec_chapter.append(max(0, len(chapters) - 1))
                parts = []
        if parts:
            warns.append(f"TOC page {p}: dropped unterminated text {' '.join(parts)!r}")
            parts = []
    return chapters, sections, sec_chapter, warns


C_TOC_HEADER = 0.075


# =============================================================================== 5. output
def hide_unused_columns(ws, used):
    """Hide every column to the right of the data, so the sheet shows only the filled columns
    (a spreadsheet grid otherwise always draws empty columns C, D, E ... next to them)."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension
    first = get_column_letter(used + 1)
    ws.column_dimensions[first] = ColumnDimension(ws, index=first, min=used + 1, max=16384, hidden=True)


def fit_rows(ws, text_col, width, line_pt=21.0):
    """Give every row the height its wrapped text needs. Without this a spreadsheet program shows long hadiths
    cut off; Excel's hard limit is 409 points per row, so the text column is made wide enough for the longest one.
    -> ids (column A) of rows that still do not fit."""
    import math
    per_line = max(20, int(width * 0.85))                 # characters that fit on one line of the column
    too_long = []
    for row in ws.iter_rows(min_row=2):
        text = row[text_col - 1].value
        if not isinstance(text, str):
            continue
        lines = sum(max(1, math.ceil(len(seg) / per_line)) for seg in text.split("\n"))
        need = lines * line_pt + 4
        if need > 409:
            too_long.append(row[0].value)
        ws.row_dimensions[row[0].row].height = min(409, need)
    return too_long


def write_workbook(path, chapters, sections, entries, font):
    """book.xlsx with three sheets -  chapters: id, name   sections: section_id, name   hadiths: id, hadith."""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows, widths in (
            ("chapters", ["id", "name"], [(i, t) for i, t in enumerate(chapters, 1)], [8, 60]),
            ("sections", ["section_id", "name"], [(i, t) for i, t in enumerate(sections, 1)], [12, 70]),
            ("hadiths", ["id", "hadith"], [(e["no"], e["bangla_text"]) for e in entries], [8, 180])):
        ws = wb.create_sheet(title)
        ws.append(header)
        for r in rows:
            ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in r])
        for row in ws.iter_rows():
            for c in row:
                head = c.row == 1
                c.font = Font(name=font, bold=head)
                c.alignment = Alignment(horizontal="center" if head else "justify" if c.column == 2 else None,
                                        vertical="top", wrap_text=True)
                if head:
                    c.fill = PatternFill("solid", fgColor="DDDDDD")
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = widths
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        hide_unused_columns(ws, 2)
        fit_rows(ws, 2, widths[1])
    wb.save(path)


def write_review(entries, path):
    """CSV of every hadith that needs a human look (UTF-8 with BOM so Excel opens Bangla correctly)."""
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["id", "pages", "ocr_conf", "flags", "text_start"])
        for e in entries:
            if e["flags"]:
                w.writerow([e["no"], f"{e['page_start']}-{e['page_end']}", e["conf_mean"],
                            "; ".join(e["flags"]), e["bangla_text"][:80]])


def assign_sections(entries, sections):
    """Index of the TOC section each hadith sits under. Headings and the TOC list are both in book order, so each
    heading change is matched against the next few sections only (a heading the OCR missed does not derail it)."""
    from difflib import SequenceMatcher
    norm = [_norm(t) for t in sections]
    cur, prev_heading, out, weak = -1, None, [], 0
    for e in entries:
        h = _norm(e["heading"])
        if e["heading"] != prev_heading:
            prev_heading = e["heading"]
            if cur == -1:
                cur = 0
            if h:
                lo = cur
                cand = [(SequenceMatcher(None, h, norm[i], autojunk=False).ratio(), i)
                        for i in range(lo, min(len(sections), lo + 7))]
                best, bi = max(cand) if cand else (0, cur)
                if best >= 0.45:
                    cur = bi
                else:                                         # not a section title (a body line taken for a heading)
                    weak += 1
        out.append(max(cur, 0))
    return out, weak


def write_docx(entries, sec_of, chap_of, chapters, sections, path, font_name, title):
    """Formatted Word file: title, chapter (Heading 1), section (Heading 2), then the numbered hadiths, justified."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)                 # A4
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    def style_font(st, size, bold=False, color=None):
        st.font.name, st.font.size, st.font.bold = font_name, Pt(size), bold
        st.font.color.rgb = color or RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(att), font_name)
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(att)) is not None:
                del rf.attrib[qn(att)]
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(size * 2)))
        rpr.append(szcs)
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "bn-BD")
        lang.set(qn("w:bidi"), "bn-BD")
        rpr.append(lang)

    normal = doc.styles["Normal"]
    style_font(normal, 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3
    style_font(doc.styles["Title"], 26, True)
    style_font(doc.styles["Heading 1"], 18, True, RGBColor(0x1F, 0x3A, 0x5F))
    style_font(doc.styles["Heading 2"], 14, True, RGBColor(0x2E, 0x2E, 0x2E))
    for name, before, after in (("Heading 1", 22, 8), ("Heading 2", 14, 4)):
        pf = doc.styles[name].paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next = Pt(before), Pt(after), True
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # centred title block
    for _ in range(6):
        doc.add_paragraph()
    doc.add_paragraph(title, style="Title")
    note = doc.add_paragraph("বাংলা অনুবাদ (OCR থেকে নিষ্কাশিত)")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # page number in the footer
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for kind, txt in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)

    doc.add_paragraph().paragraph_format.page_break_before = True      # the text starts on a new page
    last_chapter = last_section = None
    for e, si, ci in zip(entries, sec_of, chap_of):
        if ci != last_chapter and 0 <= ci < len(chapters):
            h1 = doc.add_paragraph(f"অধ্যায় : {chapters[ci]}", style="Heading 1")
            last_chapter = ci
        if si != last_section and 0 <= si < len(sections):
            doc.add_paragraph(sections[si], style="Heading 2")
            last_section = si
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r1 = para.add_run(f"[{str(e['no']).translate(TO_BN)}] " if e["no"] is not None else "")
        r1.bold = True
        para.add_run(e["bangla_text"])
    doc.core_properties.title = title
    doc.core_properties.language = "bn-BD"
    doc.save(path)


def make_book(title, chapters, sections, entries, sec_of, chap_of, sec_chapter=None):
    """Everything the outputs are made from, as one JSON-ready dict (saved as book.json; the web review edits a copy).
    section_chapter = chapter index of each section, so a hadith moved to another section follows its chapter."""
    if sec_chapter is None:                               # no TOC: a section belongs to the chapter it first appears in
        sec_chapter = [-1] * len(sections)
        for si, ci in zip(sec_of, chap_of):
            if 0 <= si < len(sections) and sec_chapter[si] == -1:
                sec_chapter[si] = ci
    hadiths = [dict(uid=f"h{i}", no=e["no"], text=e["bangla_text"], ocr_text=e["bangla_text"],
                    arabic=e.get("arabic_text", ""), foot=e.get("foot_text", ""), heading=e.get("heading", ""),
                    flags=list(e["flags"]), conf=e["conf_mean"], page_start=e["page_start"], page_end=e["page_end"],
                    anchor=e.get("anchor"), section=si, chapter=ci, reviewed=False)
               for i, (e, si, ci) in enumerate(zip(entries, sec_of, chap_of))]
    return dict(version=1, title=title, chapters=list(chapters), sections=list(sections),
                section_chapter=list(sec_chapter), hadiths=hadiths, next_uid=len(hadiths))


def export_book(book, out_dir, docx=True, font=None):
    """book.xlsx, review.csv and (unless docx=False) book.docx from a book dict. Hadiths marked reviewed are left
    out of review.csv. -> names of the files written."""
    font = font or CONFIG["font"]
    hs = book["hadiths"]
    entries = [dict(no=h["no"], bangla_text=re.sub(r"[ \t]+", " ", h["text"]).strip(), flags=[] if h.get("reviewed") else h["flags"],
                    page_start=h["page_start"], page_end=h["page_end"], conf_mean=h["conf"]) for h in hs]
    write_workbook(os.path.join(out_dir, "book.xlsx"), book["chapters"], book["sections"], entries, font)
    write_review(entries, os.path.join(out_dir, "review.csv"))
    made = ["book.xlsx", "review.csv"]
    if docx:
        write_docx(entries, [h["section"] for h in hs], [h["chapter"] for h in hs], book["chapters"],
                   book["sections"], os.path.join(out_dir, "book.docx"), font, book["title"])
        made.insert(1, "book.docx")
    return made


def write_debug(page_data, img_dir, dbg_dir):
    from PIL import Image, ImageDraw
    os.makedirs(dbg_dir, exist_ok=True)
    colors = dict(header="gray", noise="gray", arabic="red", heading="green", chapter="purple", footnote="blue",
                  body="black")
    for pd in page_data:
        img = Image.open(os.path.join(img_dir, pd["image"])).convert("RGB")
        dr = ImageDraw.Draw(img)
        for ln in pd["lines"]:
            dr.rectangle([ln["left"], ln["top"], ln["right"], ln["bottom"]], outline=colors[ln["kind"]], width=4)
            dr.text((max(0, ln["left"] - 80), ln["top"]), ln["kind"][:4], fill=colors[ln["kind"]])
        s = 1000 / img.width
        img.resize((1000, int(img.height * s))).save(os.path.join(dbg_dir, f"page_{pd['page']:03d}.jpg"), quality=70)


def pick_pdf():
    """No PDF on the command line: let the user choose one (macOS dialog, else Tk dialog, else a typed/dragged path)."""
    import subprocess
    if sys.platform == "darwin":
        r = subprocess.run(["osascript", "-e",
                            'POSIX path of (choose file of type {"pdf"} with prompt "Choose the scanned book (PDF)")'],
                           capture_output=True, text=True)
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout.strip()
        if "User canceled" in r.stderr:
            sys.exit("No PDF chosen.")
    try:
        import tkinter
        from tkinter import filedialog
        root = tkinter.Tk()
        root.withdraw()
        path = filedialog.askopenfilename(title="Choose the scanned book (PDF)", filetypes=[("PDF", "*.pdf")])
        root.destroy()
        if path:
            return path
    except Exception:
        pass
    path = input("Drag the PDF into this window and press Enter: ").strip()
    if len(path) > 1 and path[0] == path[-1] and path[0] in "'\"":
        path = path[1:-1]
    return path.replace("\\ ", " ")


def preflight(a):
    """Stop with a readable message (instead of a traceback) when something needed is not installed."""
    import importlib.util
    import shutil
    missing = [pip for mod, pip in (("pymupdf", "pymupdf"), ("pytesseract", "pytesseract"), ("PIL", "pillow"),
                                     ("openpyxl", "openpyxl"), ("docx", "python-docx"), ("numpy", "numpy"),
                                     ("cv2", "opencv-python-headless")) if importlib.util.find_spec(mod) is None]
    if missing:
        sys.exit("Missing Python packages: " + " ".join(missing) + "\n  install them with:  pip install " + " ".join(missing))
    if not (a.tesseract_cmd or shutil.which("tesseract")):
        sys.exit("Tesseract is not installed (macOS: brew install tesseract tesseract-lang; "
                 "Windows: winget install UB-Mannheim.TesseractOCR) or pass --tesseract-cmd <path to tesseract>")


# =============================================================================== main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pdf", nargs="?", help="the scanned book; leave it out to choose the file in a dialog")
    ap.add_argument("--out", help='output folder (default: "<PDF name> - output" next to the PDF)')
    ap.add_argument("--title", help="title on the first page of book.docx (default: the PDF file name)")
    ap.add_argument("--pages", help='e.g. "195-205" or "1,5,9-12" (default: all)')
    ap.add_argument("--lang", default="ara+ben")
    ap.add_argument("--workers", type=int, default=max(1, (os.cpu_count() or 2)))
    ap.add_argument("--tesseract-cmd")
    ap.add_argument("--tessdata", help="folder holding ben.traineddata / ara.traineddata")
    ap.add_argument("--debug", action="store_true", help="write annotated page images to <out>/debug")
    ap.add_argument("--keep-arabic", action="store_true",
                    help="also write the Arabic OCR into the outputs (default: dropped; the Arabic OCR is error-prone)")
    ap.add_argument("--no-docx", action="store_true", help="do not write book.docx")
    ap.add_argument("--toc-pages", default="auto",
                    help='PDF pages of the table of contents, e.g. "17-25"; "auto" (default) finds them, '
                         '"none" = the PDF has no table of contents (chapters/sections come from the headings)')
    ap.add_argument("--chapters-file", help="text file, one chapter name per line, used instead of the OCR'd "
                                            "chapter names (bold chapter lines OCR badly); default: chapters.txt "
                                            "next to this script, if present")
    ap.add_argument("--body-start", type=int, help="first PDF page with hadiths (default: page after the TOC)")
    ap.add_argument("--rebuild-only", action="store_true", help="skip image extraction and OCR; reuse cached JSON")
    a = ap.parse_args()
    for stream in (sys.stdout, sys.stderr):               # the Windows console is not UTF-8: Bangla must not crash it
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    if not a.tesseract_cmd and sys.platform == "win32":   # the Windows installer does not put tesseract on PATH
        import shutil
        if not shutil.which("tesseract"):
            for base in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)"),
                         os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs")):
                exe = os.path.join(base or "", "Tesseract-OCR", "tesseract.exe")
                if os.path.isfile(exe):
                    a.tesseract_cmd = exe
                    break
    preflight(a)
    picked = not a.pdf
    if picked:
        a.pdf = pick_pdf()
    a.pdf = os.path.abspath(os.path.expanduser(a.pdf))
    if not os.path.isfile(a.pdf):
        sys.exit(f"PDF not found: {a.pdf}")
    if not a.out:                                         # one folder per PDF: the page cache must never be shared
        a.out = os.path.join(os.path.dirname(a.pdf), os.path.splitext(os.path.basename(a.pdf))[0] + " - output")
    print(f"PDF:    {a.pdf}\nOutput: {a.out}")
    os.makedirs(a.out, exist_ok=True)
    stamp, me = os.path.join(a.out, "source.json"), dict(name=os.path.basename(a.pdf), size=os.path.getsize(a.pdf))
    if os.path.exists(stamp):
        if json.load(open(stamp, encoding="utf8")) != me:
            sys.exit(f"{a.out}\n  holds the results of a different PDF. Use another --out folder, or delete this one.")
    elif os.path.isdir(os.path.join(a.out, "ocr")) and os.listdir(os.path.join(a.out, "ocr")):
        sys.exit(f"{a.out}\n  already holds cached pages of unknown origin. Use another --out folder, or delete this one.")
    else:
        json.dump(me, open(stamp, "w", encoding="utf8"), ensure_ascii=False)

    here = os.path.dirname(os.path.abspath(__file__))
    if not a.tessdata and os.path.exists(os.path.join(here, "tessdata_best", "ben.traineddata")):
        a.tessdata = os.path.join(here, "tessdata_best")             # more accurate models, used automatically
    print(f"Tesseract models: {a.tessdata or 'system default (fast models)'}")
    img_dir, ocr_dir = os.path.join(a.out, "images"), os.path.join(a.out, "ocr")
    os.makedirs(a.out, exist_ok=True)

    import pymupdf
    n = len(pymupdf.open(a.pdf))
    if a.toc_pages == "auto":
        toc_pages = detect_toc_pages(a.pdf, img_dir, a.out, n)
    else:
        toc_pages = [] if a.toc_pages == "none" else parse_pages(a.toc_pages, n)
    body_start = a.body_start or (max(toc_pages) + 1 if toc_pages else 1)

    if not a.rebuild_only:
        pages = parse_pages(a.pages, n)
        print(f"{n} pages in PDF, processing {len(pages)}")
        paths = extract_images(a.pdf, img_dir, pages)
        run_ocr(paths, ocr_dir, a)
        wanted = set(pages)
    else:
        wanted = parse_pages(a.pages, 10 ** 6) if a.pages else None

    page_data = []
    for f in sorted(os.listdir(ocr_dir)):
        if f.endswith(".json"):
            d = json.load(open(os.path.join(ocr_dir, f), encoding="utf8"))
            if (wanted is None or d["page"] in wanted) and d["page"] >= body_start:
                page_data.append(d)
    page_data.sort(key=lambda d: d["page"])
    for d in page_data:
        d["rule_y"] = find_rule(os.path.join(img_dir, d["image"]))
    cf = a.chapters_file or os.path.join(here, "chapters.txt")
    known = [l.strip() for l in open(cf, encoding="utf8") if l.strip()] if os.path.exists(cf) else []
    warns = []
    for d in page_data:
        classify_page(d, CONFIG)
    sec_chapter = None
    if toc_pages:
        print(f"Table of contents: pages {toc_pages[0]}-{toc_pages[-1]}")
        chapters, sections, sec_chapter, warns = read_toc(img_dir, toc_pages, a.out, a)
        print(f"Headings confirmed against the TOC: {apply_toc_headings(page_data, sections, CONFIG)}")
        structural_headings(page_data, CONFIG)
        mark_chapters(page_data, img_dir, a, known, reread=False)
        entries, unnumbered = build_entries(page_data, CONFIG)
        repair_numbers(entries)
        if known and (len(known) == len(chapters) or a.chapters_file):
            print(f"Chapter names taken from {cf} ({len(known)})")
            chapters = known
        elif known:                                           # chapters.txt belongs to another book
            warns.append(f"{os.path.basename(cf)} has {len(known)} names but this book has {len(chapters)} "
                         f"chapters: ignored, OCR'd chapter names used")
        sec_of, weak = assign_sections(entries, sections)
        if weak:
            print(f"{weak} heading-like line(s) matched no TOC section and were ignored")
        chap_of = [sec_chapter[si] if si < len(sec_chapter) else -1 for si in sec_of]
    else:
        print("No table of contents in this PDF: chapters and sections are taken from the headings in the text")
        structural_headings(page_data, CONFIG)
        mark_chapters(page_data, img_dir, a, known, reread=True)
        reread_headings(page_data, img_dir, a)
        for d in page_data:
            for ln in d["lines"]:
                if ln.get("chapter_unsure"):
                    warns.append(f"page {d['page']}: chapter title read as {ln['chapter_name']!r} - check it")
        entries, unnumbered = build_entries(page_data, CONFIG)
        repair_numbers(entries)
        chapters, sections, chap_of, sec_of = structure_from_headings(entries)
    title = a.title or running_title(page_data) or os.path.splitext(os.path.basename(a.pdf))[0]
    book = make_book(title, chapters, sections, entries, sec_of, chap_of, sec_chapter)
    with open(os.path.join(a.out, "book.json"), "w", encoding="utf8") as fh:
        json.dump(book, fh, ensure_ascii=False)
    made = export_book(book, a.out, docx=not a.no_docx)
    for old in ("chapters.xlsx", "sections.xlsx", "hadith.xlsx"):      # files of earlier versions: now sheets of book.xlsx
        if os.path.exists(os.path.join(a.out, old)):
            os.remove(os.path.join(a.out, old))
    if a.debug:
        write_debug(page_data, img_dir, os.path.join(a.out, "debug"))

    flagged = sum(1 for e in entries if e["flags"])
    print(f"\nDone. {len(chapters)} chapters, {len(sections)} sections, {len(entries)} hadiths "
          f"({flagged} need review -> review.csv), {len(unnumbered)} orphan lines.")
    for w in warns:
        print("  warning:", w)
    print("Created:")
    for name in made:
        print("  " + os.path.join(a.out, name))
    if picked and sys.platform == "darwin":               # chosen in a dialog -> show the results folder
        import subprocess
        subprocess.run(["open", a.out])
    elif picked and sys.platform == "win32":
        os.startfile(a.out)
    if a.debug:
        print(f"Debug images: {os.path.join(a.out, 'debug')}  (red=Arabic, green=heading, blue=footnote, gray=dropped)")


if __name__ == "__main__":
    main()
